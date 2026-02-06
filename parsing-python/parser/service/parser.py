import pandas as pd
import openpyxl as xl
import json
import re

from service.Lesson import Lesson

def __find_head_position(df): # остаётся молиться, что в xlsx первая нужная строка будет начинаться с "Дни"
    name_first_col = "Дни"
    for i in range(df.shape[0]):
        if df.iloc[i].values[0] == name_first_col: return i

def __create_list_delete_rows(head_position): # решил вынести логику для формирования списка ненужных строк
    delete_position = []
    for i in range(head_position):
        delete_position.append(i)
    return delete_position

def __get_group_column(df): # сразу определяю индексы групп
    groups = []
    no_groups_word = frozenset({"Дни", "Часы", None})
    for i in range(df.shape[1]):
        buffer = df.iloc[0].values[i]
        if buffer not in no_groups_word: groups.append(i)
    return groups
    
def __extract_lesson(df, row_index, col_index): # выношу логику сбора данных о паре
    days_of_week = frozenset({"ПОНЕДЕЛЬНИК", "ВТОРНИК", "СРЕДА", "ЧЕТВЕРГ", "ПЯТНИЦА", "СУББОТА", "ВОСКРЕСЕНЬЕ"})
    day_of_week = str(df.iat[row_index, 0]).strip()
    if day_of_week in days_of_week:  # исключаю нижние строки для подписей
        data = str(df.iat[row_index, col_index]).strip()
        if not data or data == "nan" or len(data) == 0: return None # чтобы не записывал пустые пары
        type, subject, teachers, rooms = __extract_data(data)

        group = df.iat[0, col_index].strip()

        time = df.iat[row_index, 1].strip().split('-')
        start_time = time[0]
        end_time = time[1]

        even = row_index % 2 == 0
        lesson = Lesson(group, day_of_week, start_time, end_time, type, subject, teachers, rooms, even)
        return lesson
    else: 
        return None
    
def __extract_data(data):
    type_pattern = r'^([^.]+)\.'
    type_re = re.compile(type_pattern)
    teacher_pattern = r'([А-ЯЁ][а-яё\-]+(?:\s+[А-ЯЁ][а-яё\-]+)?\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)'
    teacher_re = re.compile(teacher_pattern)
    room_pattern = r'\b\d{1,2}-[А-ЯA-Z0-9Ёё\- ]+(?:\s*\([^)]+\))?|\bСК\b|\bКСК-\d+\b'
    room_re = re.compile(room_pattern)

    cleaned_data = str(data).replace('\xa0', ' ').replace('\n', ' ').replace('\t', ' ').strip()

    lesson_type = None
    m = type_re.match(cleaned_data)
    if m:
        lesson_type = m.group(1)
        cleaned_data = cleaned_data[m.end():].strip()

    teacher_matches = list(teacher_re.finditer(cleaned_data))
    teachers = [m.group(1) for m in teacher_matches]

    room_matches = list(room_re.finditer(cleaned_data))
    rooms = [m.group(0).strip() for m in room_matches]

    cut_positions = []

    if teacher_matches:
        cut_positions.append(teacher_matches[0].start())

    if room_matches:
        cut_positions.append(room_matches[0].start())

    subject_end = min(cut_positions) if cut_positions else len(cleaned_data)
    subject = cleaned_data[:subject_end].strip()

    return lesson_type, subject, teachers, rooms

def __handle_merge_cells(work_sheet):
    merged_cells = list(work_sheet.merged_cells.ranges)

    for merge_cell in merged_cells:
        work_sheet.unmerge_cells(str(merge_cell))
        statColumn = merge_cell.min_col
        top_left_cell = work_sheet.cell(row = merge_cell.min_row, column = statColumn)
        for row in range(merge_cell.min_row, merge_cell.max_row + 1):
            empty_cell = work_sheet.cell(row = row, column=statColumn)
            empty_cell.value = top_left_cell.value

    return work_sheet

def __extract_shedule(data_frame):
    index_group = __get_group_column(data_frame)
    shedule = []
    for col in index_group:
        for row in range(1, data_frame.shape[0]):
            lesson = __extract_lesson(data_frame, row, col)
            if lesson != None: shedule.append(lesson)
    return shedule

def __serialize_shedule_to_json(shedule):
    new_shedule = [lesson.__dict__ for lesson in shedule]
    file_path = './resource/schedule.json'
    with open(file_path, 'w', encoding = "UTF-8") as outfile:
        json.dump(new_shedule, outfile, indent = 4, ensure_ascii = False)

    return file_path

def parsing(file): 
    book = xl.load_workbook(file, data_only = False, read_only = False)
    shedule = []
    for sheet in book:
        handled_sheet = __handle_merge_cells(sheet)
        df = pd.DataFrame(handled_sheet.values)
        head_positions = __find_head_position(df)
        df.drop(index = __create_list_delete_rows(head_positions), inplace = True) 
        shedule.extend(__extract_shedule(df))

    return __serialize_shedule_to_json(shedule)


