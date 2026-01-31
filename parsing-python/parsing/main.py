import pandas as pd
import openpyxl as xl

from Lesson import Lesson

def find_head_position(df): # остаётся молиться, что в xlsx первая нужная строка будет начинаться с "Дни"
    name_first_col = "Дни"
    for i in range(df.shape[0]):
        if df.iloc[i].values[0] == name_first_col: return i

def create_list_delete_rows(head_position): # решил вынести логику для формирования списка ненужных строк
    delete_position = []
    for i in range(head_position):
        delete_position.append(i)
    return delete_position

def get_group_column(df): # сразу определяю индексы групп
    groups = []
    no_groups_word = frozenset({"Дни", "Часы", None})
    for i in range(df.shape[1]):
        buffer = df.iloc[0].values[i]
        if buffer not in no_groups_word: groups.append(i)
    return groups
    
def extract_lesson(df, row_index, col_index): # выношу логику сбора данных о паре
    days_of_week = frozenset({"ПОНЕДЕЛЬНИК", "ВТОРНИК", "СРЕДА", "ЧЕТВЕРГ", "ПЯТНИЦА", "СУББОТА", "ВОСКРЕСЕНЬЕ"})
    day_of_week = str(df.iat[row_index, 0]).strip()
    if day_of_week in days_of_week:  # исключаю нижние строки для подписей
        data = str(df.iat[row_index, col_index]).strip()
        if day_of_week == "СРЕДА": print(day_of_week)
        if len(data) == 0: return None # чтобы не записывал пустые пары
        group = df.iat[0, col_index].strip()
        time = df.iat[row_index, 1].strip()
        even = row_index % 2 == 0
        lesson = Lesson(group, day_of_week, time, data, even)
        return lesson
    else: 
        return None


book = xl.load_workbook('./resource/shedule.xlsx', data_only = False, read_only = False)
sheet = book.active

merged_cells = list(sheet.merged_cells.ranges)

# цикл для разъединения объекдинённых клеток и заполнения, ставших пустыми, клетки 
# вынести в метод
for merge_cell in merged_cells:
    sheet.unmerge_cells(str(merge_cell))
    statColumn = merge_cell.min_col
    top_left_cell = sheet.cell(row = merge_cell.min_row, column = statColumn)
    for row in range(merge_cell.min_row, merge_cell.max_row + 1):
        empty_cell = sheet.cell(row = row, column=statColumn)
        empty_cell.value = top_left_cell.value

df = pd.DataFrame(sheet.values)
head_positions = find_head_position(df)
# head = df.iloc(headPosition) если нужно будет установить заголовок вместо цифр
df.drop(index = create_list_delete_rows(head_positions), inplace = True)
index_group = get_group_column(df)

shedule = []

for col in index_group:
    for row in range(1, df.shape[0]):
        lesson = extract_lesson(df, row, col)
        if lesson != None: shedule.append(lesson)

for lesson in shedule: 
    lesson.describe()

print(len(shedule))




